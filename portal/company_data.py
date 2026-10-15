"""Synthetic but realistic document estate for the fictional company *Halvorsen Industrial Systems*.

Everything is generated from one seeded "world" (people, departments, sites, customers, suppliers,
products) so documents are mutually consistent: the account manager named in the sales pipeline
exists in the staff directory, budget owners are real department heads, and so on.

Each document carries access metadata used by the portal:
    level 1  All staff            - everyone
    level 2  Department           - members of the listed departments (+ managers and executives)
    level 3  Management           - managers and executives
    level 4  Executive / Board    - executives only

All names, companies, figures and contact details are fictional (Faker, fixed seed).
"""
from __future__ import annotations

import csv
import json
import random
import shutil
from dataclasses import asdict, dataclass, field
from datetime import date, timedelta
from pathlib import Path

from faker import Faker

from rag import config
from rag.corpus import COMPANY, SITES, _pdf_doc

PORTAL_DIR = config.DATA_DIR / "portal"
DOC_DIR = PORTAL_DIR / "documents"
UPLOAD_DIR = PORTAL_DIR / "uploads"
MANIFEST = PORTAL_DIR / "manifest.json"
ACCOUNTS = PORTAL_DIR / "demo_accounts.csv"

DEPARTMENTS = ["Executive Office", "Finance", "Human Resources", "Engineering", "Operations", "Sales",
               "Marketing", "Legal & Compliance", "IT", "Quality", "Procurement"]
LEVELS = {1: "All staff", 2: "Department", 3: "Management", 4: "Executive"}
SITE_NAMES = [s for s, _ in SITES]
SITE_COUNTRY = dict(SITES)
LOCALES = {"Finland": "fi_FI", "Poland": "pl_PL", "Portugal": "pt_PT", "Czechia": "cs_CZ", "France": "fr_FR",
           "Spain": "es_ES", "Austria": "de_AT", "Denmark": "da_DK"}
PRODUCTS = [f"HX-{n}" for n in (110, 140, 160, 200, 250, 280, 320, 360, 400, 450, 520, 600)] + \
           ["CU-4 Controller", "CU-4 Pro Controller", "SmartFlow Monitoring (SaaS)", "Service Agreement Gold",
            "Service Agreement Silver", "Spare Parts Kit"]


@dataclass
class DocMeta:
    doc: str
    title: str
    departments: list[str]
    level: int
    owner: str
    created: str
    description: str
    category: str
    source: str = "generated"
    uploaded_by: str | None = None


@dataclass
class Person:
    emp_id: str
    name: str
    title: str
    department: str
    site: str
    email: str
    phone: str
    manager: str
    start: date
    grade: str
    salary: int = 0


@dataclass
class World:
    rng: random.Random
    fake: Faker
    execs: dict[str, Person] = field(default_factory=dict)
    heads: dict[str, Person] = field(default_factory=dict)
    staff: list[Person] = field(default_factory=list)
    customers: list[dict] = field(default_factory=list)
    suppliers: list[str] = field(default_factory=list)
    docs: list[DocMeta] = field(default_factory=list)
    fin: list[dict] = field(default_factory=list)      # one consistent set of quarterly financials (EUR m)

    def person(self, dept: str) -> Person:
        return self.rng.choice([p for p in self.staff if p.department == dept])


# ---------------------------------------------------------------------------
# world building
# ---------------------------------------------------------------------------
EXEC_ROLES = [("ceo", "Chief Executive Officer", "Executive Office"), ("cfo", "Chief Financial Officer", "Finance"),
              ("coo", "Chief Operating Officer", "Operations"), ("cto", "Chief Technology Officer", "Engineering"),
              ("chro", "Chief Human Resources Officer", "Human Resources"),
              ("cco", "Chief Commercial Officer", "Sales"), ("cio", "Chief Information Officer", "IT"),
              ("gc", "General Counsel", "Legal & Compliance")]
TITLES = {
    "Finance": ["Financial Controller", "Senior Accountant", "Accountant", "FP&A Analyst", "Treasury Analyst", "Payroll Specialist", "AP Clerk"],
    "Human Resources": ["HR Business Partner", "Recruiter", "HR Generalist", "Learning & Development Specialist", "Compensation Analyst"],
    "Engineering": ["Principal Engineer", "Senior Mechanical Engineer", "Mechanical Engineer", "Firmware Engineer", "Test Engineer",
                    "CFD Analyst", "Product Manager", "Electrical Engineer", "Reliability Engineer"],
    "Operations": ["Production Supervisor", "Shift Lead", "Maintenance Technician", "Planner", "Process Engineer",
                   "Logistics Coordinator", "Machine Operator", "EHS Specialist"],
    "Sales": ["Key Account Manager", "Account Manager", "Sales Engineer", "Inside Sales Representative", "Sales Operations Analyst"],
    "Marketing": ["Marketing Manager", "Product Marketing Specialist", "Digital Marketing Specialist", "Content Writer"],
    "Legal & Compliance": ["Senior Legal Counsel", "Contract Manager", "Compliance Officer", "Paralegal"],
    "IT": ["Systems Administrator", "Network Engineer", "Service Desk Analyst", "Security Analyst", "ERP Consultant"],
    "Quality": ["Quality Engineer", "Quality Inspector", "Supplier Quality Engineer", "Metrology Technician"],
    "Procurement": ["Category Buyer", "Strategic Sourcing Manager", "Purchasing Assistant"],
    "Executive Office": ["Executive Assistant", "Strategy Analyst", "Chief of Staff"],
}
DEPT_WEIGHT = {"Operations": 30, "Engineering": 16, "Sales": 10, "Quality": 8, "Finance": 6, "IT": 6,
               "Human Resources": 5, "Procurement": 5, "Marketing": 4, "Legal & Compliance": 3, "Executive Office": 2}


def _email(name: str) -> str:
    import unicodedata
    n = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode().lower()
    return ".".join(p for p in n.replace("'", "").split() if p) + "@halvorsen-is.example"


def _new_person(w: World, dept: str, title: str, site: str, manager: str, idx: int, grade: str) -> Person:
    loc = LOCALES[SITE_COUNTRY[site]]
    f = Faker(loc)
    f.seed_instance(w.rng.randint(0, 10**9))
    name = f"{f.first_name()} {f.last_name()}"
    start = date(2004, 1, 1) + timedelta(days=w.rng.randrange(0, 7900))
    cc = {"Finland": "358", "Poland": "48", "Portugal": "351", "Czechia": "420", "France": "33", "Spain": "34", "Austria": "43", "Denmark": "45"}
    phone = f"+{cc[SITE_COUNTRY[site]]} {w.rng.randint(10, 99)} {w.rng.randint(100, 999)} {w.rng.randint(1000, 9999)}"
    return Person(f"E{10000 + idx}", name, title, dept, site, _email(name), phone, manager, start, grade)


def build_world(seed: int = 2026, n_staff: int = 650) -> World:
    rng = random.Random(seed)
    fake = Faker(["en_GB", "de_DE", "fi_FI", "fr_FR"])
    Faker.seed(seed)
    w = World(rng, fake)
    idx = 0
    for key, title, dept in EXEC_ROLES:
        p = _new_person(w, dept, title, "Tampere Works", "Board of Directors" if key == "ceo" else "", idx, "E1")
        idx += 1
        w.execs[key] = p
    for p in w.execs.values():
        if not p.manager:
            p.manager = w.execs["ceo"].name
    exec_of = {p.department: p for k, p in w.execs.items() if k != "ceo"}
    for dept in DEPARTMENTS:
        boss = exec_of.get(dept, w.execs["ceo"])
        title = "Chief of Staff" if dept == "Executive Office" else f"Head of {dept}"
        site = "Tampere Works" if dept in ("Executive Office", "Finance", "Legal & Compliance") else rng.choice(SITE_NAMES)
        h = _new_person(w, dept, title, site, boss.name, idx, "M1")
        idx += 1
        w.heads[dept] = h
    w.staff = list(w.execs.values()) + list(w.heads.values())
    depts = list(DEPT_WEIGHT)
    weights = [DEPT_WEIGHT[d] for d in depts]
    while len(w.staff) < n_staff:
        dept = rng.choices(depts, weights)[0]
        title = rng.choice(TITLES[dept])
        grade = "P3" if any(t in title for t in ("Principal", "Senior", "Key", "Strategic", "Controller")) else rng.choice(["P1", "P2", "P2", "S2", "S3"])
        site = rng.choice(SITE_NAMES) if dept in ("Operations", "Quality", "Engineering", "Sales") else rng.choice(SITE_NAMES[:3] + ["Tampere Works"] * 3)
        w.staff.append(_new_person(w, dept, title, site, w.heads[dept].name, len(w.staff), grade))
    names = set()
    for p in w.staff:            # de-duplicate names so look-ups are unambiguous
        while p.name in names:
            p.name = p.name + " Jr."
        names.add(p.name)
        p.email = _email(p.name)
    bands = {"E1": (210_000, 420_000), "M1": (125_000, 190_000), "P3": (78_000, 118_000), "P2": (56_000, 82_000),
             "P1": (42_000, 60_000), "S3": (38_000, 52_000), "S2": (31_000, 42_000)}
    for p in w.staff:
        lo, hi = bands[p.grade]
        p.salary = rng.randrange(lo, hi, 500)

    sectors = [("Water Utilities", ["Water", "Water Services", "Waterworks"]), ("District Energy", ["Energy", "District Heating", "Power & Heat"]),
               ("Chemicals", ["Chemicals", "Specialty Chemicals", "Process Chemicals"]), ("Food & Beverage", ["Foods", "Dairy", "Brewery", "Beverages"]),
               ("Mining", ["Mining", "Minerals", "Resources"]), ("HVAC Contractors", ["Building Services", "Climate Systems", "Installations"]),
               ("Pulp & Paper", ["Paper Mills", "Pulp & Board", "Packaging"])]
    legal = {"Finland": ["Oy", "Oyj"], "Germany": ["GmbH", "AG"], "Portugal": ["S.A.", "Lda."], "France": ["SAS", "SA"], "Spain": ["S.A.", "S.L."],
             "Austria": ["GmbH", "AG"], "Denmark": ["A/S", "ApS"], "Poland": ["Sp. z o.o.", "S.A."], "Czechia": ["s.r.o.", "a.s."],
             "Sweden": ["AB"], "Estonia": ["AS", "OÜ"]}
    cfaker = {c: Faker(l) for c, l in {"Finland": "fi_FI", "Germany": "de_DE", "Portugal": "pt_PT", "France": "fr_FR", "Spain": "es_ES",
                                        "Austria": "de_AT", "Denmark": "da_DK", "Poland": "pl_PL", "Czechia": "cs_CZ", "Sweden": "sv_SE", "Estonia": "et_EE"}.items()}
    for f_ in cfaker.values():
        f_.seed_instance(seed)
    cities = ["Oulu", "Lahti", "Bremen", "Leipzig", "Porto", "Braga", "Lyon", "Nantes", "Bilbao", "Valencia", "Graz",
              "Linz", "Aarhus", "Odense", "Gdańsk", "Poznań", "Brno", "Ostrava", "Uppsala", "Malmö", "Turku", "Tartu"]
    countries = {"Oulu": "Finland", "Lahti": "Finland", "Turku": "Finland", "Bremen": "Germany", "Leipzig": "Germany",
                 "Porto": "Portugal", "Braga": "Portugal", "Lyon": "France", "Nantes": "France", "Bilbao": "Spain",
                 "Valencia": "Spain", "Graz": "Austria", "Linz": "Austria", "Aarhus": "Denmark", "Odense": "Denmark",
                 "Gdańsk": "Poland", "Poznań": "Poland", "Brno": "Czechia", "Ostrava": "Czechia", "Uppsala": "Sweden",
                 "Malmö": "Sweden", "Tartu": "Estonia"}
    sales_people = [p for p in w.staff if p.department == "Sales" and "Manager" in p.title] or [p for p in w.staff if p.department == "Sales"]
    used = set()
    while len(w.customers) < 320:
        sector, words = rng.choice(sectors)
        city = rng.choice(cities)
        country = countries[city]
        stem = city if rng.random() < .55 else cfaker[country].last_name()
        name = f"{stem} {rng.choice(words)} {rng.choice(legal[country])}"
        if name in used:
            continue
        used.add(name)
        w.customers.append({"account_id": f"C-{20000 + len(w.customers)}", "account_name": name, "country": countries[city],
                            "city": city, "industry": sector, "segment": rng.choices(["Strategic", "Key", "Core", "Transactional"], [1, 3, 6, 8])[0],
                            "annual_revenue_eur": rng.randrange(18_000, 2_400_000, 500), "account_manager": rng.choice(sales_people).name,
                            "customer_since": rng.randint(1998, 2025), "payment_terms_days": rng.choice([30, 30, 45, 60, 90]),
                            "credit_limit_eur": rng.choice([50_000, 100_000, 250_000, 500_000, 1_000_000])})
    sup_used = set()
    a = ["Brightline", "Nordwave", "Ironpeak", "Silverfjord", "Bluecrest", "Redstone", "Clearwater", "Polaris", "Granite", "Evergreen",
         "Northgate", "Summit", "Keystone", "Harbour", "Falcon", "Meridian", "Oakridge", "Vantage", "Coastal", "Alpine", "Baltic", "Danube"]
    b = ["Castings", "Fasteners", "Seals", "Bearings", "Electronics", "Motors", "Polymers", "Steelworks", "Coatings", "Logistics", "Machining", "Cables"]
    while len(w.suppliers) < 120:
        s = f"{rng.choice(a)} {rng.choice(b)}"
        if s not in sup_used:
            sup_used.add(s)
            w.suppliers.append(s)
    w.fin = _financials(rng)
    return w


def _financials(rng: random.Random, fy_revenue: float = 412.6) -> list[dict]:
    """Quarterly FY2025 figures that add up to the published annual revenue (EUR 412.6m)."""
    shares = [rng.uniform(.235, .255) for _ in range(3)]
    shares.append(1 - sum(shares))
    out = []
    for q, sh in enumerate(shares, 1):
        rev = round(fy_revenue * sh, 1)
        cogs = round(rev * rng.uniform(.62, .655), 1)
        opex = round(rev * rng.uniform(.205, .225), 1)
        ebitda = round(rev - cogs - opex, 1)
        da = round(rng.uniform(3.2, 3.9), 1)
        ni = round((ebitda - da - rng.uniform(.4, .8)) * .8, 1)
        oi = round(rev * rng.uniform(1.02, 1.16), 1)
        out.append({"quarter": f"Q{q} FY2025", "revenue": rev, "budget_revenue": round(rev * rng.uniform(.96, 1.03), 1),
                    "cogs": cogs, "gross_margin_pct": round(100 * (rev - cogs) / rev, 1), "opex": opex, "ebitda": ebitda,
                    "ebitda_margin_pct": round(100 * ebitda / rev, 1), "depreciation": da, "net_income": ni, "order_intake": oi,
                    "book_to_bill": round(oi / rev, 2), "free_cash_flow": round(rng.uniform(4, 12), 1), "cash": round(rng.uniform(30, 60), 1),
                    "receivables": round(rng.uniform(55, 70), 1), "inventories": round(rng.uniform(48, 62), 1), "payables": round(rng.uniform(35, 45), 1),
                    "net_debt": round(rng.uniform(-20, 15), 1), "equity": round(rng.uniform(190, 215), 1), "dso_days": rng.randint(48, 61),
                    "ltifr": round(rng.uniform(.45, .8), 2), "on_time_delivery_pct": round(rng.uniform(89, 95.5), 1)})
    return out


# ---------------------------------------------------------------------------
# writers
# ---------------------------------------------------------------------------
def _add(w: World, doc, title, depts, level, owner: Person, desc, category, created: date | None = None):
    created = created or date(2026, 1, 5) + timedelta(days=w.rng.randrange(0, 260))
    w.docs.append(DocMeta(doc, title, depts, level, owner.name, created.isoformat(), desc, category))


def _csv(w: World, name, header, rows, title, depts, level, owner, desc, category):
    with open(DOC_DIR / name, "w", newline="") as f:
        wr = csv.writer(f)
        wr.writerow(header)
        wr.writerows(rows)
    _add(w, name, title, depts, level, owner, desc, category)


def _md(w: World, name, text, title, depts, level, owner, desc, category):
    (DOC_DIR / name).write_text(text)
    _add(w, name, title, depts, level, owner, desc, category)


def _pdf(w: World, name, title, subtitle, pages, depts, level, owner, desc, category, created: date | None = None):
    _pdf_doc(DOC_DIR / name, title, subtitle, pages)
    _add(w, name, title, depts, level, owner, desc, category, created)


def _docx(w: World, name, title, sections, depts, level, owner, desc, category):
    """sections: list of (heading, [paragraphs], table rows | None); one section per page."""
    from docx import Document

    d = Document()
    d.add_heading(title, 0)
    d.add_paragraph(f"{COMPANY} - {LEVELS[level]} - owner: {owner.name}")
    for i, (head, paras, table) in enumerate(sections):
        if i:
            d.add_page_break()
        d.add_heading(head, level=1)
        for p in paras:
            d.add_paragraph(p)
        if table:
            t = d.add_table(rows=0, cols=len(table[0]))
            t.style = "Table Grid"
            for row in table:
                cells = t.add_row().cells
                for c, v in zip(cells, row):
                    c.text = str(v)
    d.save(DOC_DIR / name)
    _add(w, name, title, depts, level, owner, desc, category)


def eur(x: float) -> str:
    return f"EUR {x:,.0f}"


def m(x: float) -> str:
    return f"EUR {x / 1e6:,.1f}m"


# ---------------------------------------------------------------------------
# document builders, by department
# ---------------------------------------------------------------------------
def executive_docs(w: World):
    r, ceo, cfo = w.rng, w.execs["ceo"], w.execs["cfo"]
    board = [ceo.name, "Margareta Lindqvist (Chair)", "Henrik Aalto", "Dr. Ines Carvalho", "Pieter van Dam", cfo.name]
    decisions = [
        "approved capital expenditure of {cap} for the Porto Foundry induction furnace upgrade",
        "approved an interim dividend of EUR {div} per share",
        "authorised management to begin exclusive due diligence on Kestrel Fluid Controls (Project Nordlys)",
        "approved the FY2026 operating budget of {bud}",
        "approved the opening of a sales office in Stockholm in Q3 2026",
        "noted the ISO 9001 surveillance audit passed with zero major findings",
        "approved a share buy-back programme of up to {buy}",
        "approved the appointment of Nordic Assurance Services as external quality auditor for a further three years",
        "requested a cybersecurity maturity assessment following the identity provider outage",
        "approved the 2026 long-term incentive plan for the executive team"]
    for q, d in enumerate([date(2025, 4, 23), date(2025, 7, 22), date(2025, 10, 21), date(2026, 1, 27)], start=1):
        f = w.fin[q - 1]
        rev = f["revenue"]
        pages = [{"heading": "1. Attendance and quorum", "paragraphs": [
            f"Meeting held on {d.strftime('%d %B %Y')} at Tampere Works, Finland. Present: {', '.join(board)}. "
            f"Also in attendance: {w.execs['coo'].name} (COO) and {w.execs['gc'].name} (General Counsel, secretary to the board). Quorum was confirmed.",
            f"The minutes of the previous meeting were approved without amendment."]},
            {"heading": "2. Financial performance", "paragraphs": [
                f"The CFO reported Q{q} FY2025 revenue of EUR {rev:.1f} million against a budget of EUR {f['budget_revenue']:.1f} million, "
                f"with an adjusted EBITDA margin of {f['ebitda_margin_pct']:.1f}% and cash of EUR {f['cash']:.1f} million.",
                f"Order intake for the quarter was EUR {f['order_intake']:.1f} million, giving a book-to-bill ratio of {f['book_to_bill']:.2f}."],
             "table": [["KPI", f"Q{q} FY2025", "Budget"], ["Revenue", f"EUR {rev:.1f}m", f"EUR {f['budget_revenue']:.1f}m"],
                       ["EBITDA margin", f"{f['ebitda_margin_pct']:.1f}%", "14.5%"], ["LTIFR", f"{f['ltifr']:.2f}", "0.70"],
                       ["On-time delivery", f"{f['on_time_delivery_pct']:.1f}%", "93.0%"]]},
            {"heading": "3. Resolutions", "paragraphs": [
                "The Board " + decisions[(q * 2) % len(decisions)].format(cap=m(r.uniform(4e6, 12e6)), div=f"{r.uniform(.2, .6):.2f}",
                                                                          bud=m(r.uniform(380e6, 420e6)), buy=m(r.uniform(5e6, 15e6))) + ".",
                "The Board " + decisions[(q * 2 + 1) % len(decisions)].format(cap=m(r.uniform(4e6, 12e6)), div=f"{r.uniform(.2, .6):.2f}",
                                                                              bud=m(r.uniform(380e6, 420e6)), buy=m(r.uniform(5e6, 15e6))) + ".",
                f"The next meeting is scheduled for {(d + timedelta(days=91)).strftime('%d %B %Y')}."]}]
        _pdf(w, f"Board_Minutes_2025_Q{q}.pdf", f"Board of Directors - Minutes Q{q} 2025", "Strictly confidential", pages,
             ["Executive Office"], 4, w.execs["gc"], f"Minutes of the Q{q} 2025 board meeting: performance, resolutions and approvals.", "Board",
             d + timedelta(days=7))

    targets = [("2026", 438, 15.2, 4050), ("2027", 476, 16.0, 4180), ("2028", 521, 17.1, 4300)]
    pages = [{"heading": "1. Vision", "paragraphs": [
        f"By 2028 {COMPANY} will be the most reliable European supplier of engineered pumping systems, with recurring "
        "service and digital revenue representing 25% of group sales."]},
             {"heading": "2. Financial targets", "paragraphs": ["Targets approved by the Board in November 2025."],
              "table": [["Year", "Revenue", "Adj. EBITDA margin", "Headcount"]] + [[y, f"EUR {rv}m", f"{mg}%", f"{h:,}"] for y, rv, mg, h in targets]},
             {"heading": "3. Strategic initiatives", "paragraphs": [
                 f"Initiative 1 - SmartFlow digital services: owner {w.execs['cto'].name}, budget EUR 9.5m, target 4,000 connected pumps by 2027.",
                 f"Initiative 2 - Footprint optimisation: owner {w.execs['coo'].name}, consolidate machining from Brno into Gdansk Assembly by Q2 2027, saving EUR 6.1m per year.",
                 f"Initiative 3 - Nordic growth: owner {w.execs['cco'].name}, open Stockholm sales office and grow Nordic revenue by 40%.",
                 f"Initiative 4 - Project Nordlys: owner {ceo.name}, bolt-on acquisition in valve automation (see confidential memo).",
                 f"Initiative 5 - Talent: owner {w.execs['chro'].name}, reduce voluntary attrition from 9.8% to below 7%."]},
             {"heading": "4. Risks", "paragraphs": [
                 "Principal risks: energy price volatility, dependence on two casting suppliers, cyber risk to OT networks, and "
                 "integration risk on acquisitions. Mitigations are owned by the executive team and reviewed quarterly."]}]
    _pdf(w, "Strategic_Plan_2026-2028.pdf", "Strategic Plan 2026-2028", "Board approved - confidential", pages,
         ["Executive Office"], 4, ceo, "Three-year strategy: vision, financial targets, five initiatives and key risks.", "Strategy", date(2025, 11, 28))

    ev = r.uniform(58, 74)
    pages = [{"heading": "1. Transaction overview", "paragraphs": [
        "Project Nordlys is the proposed acquisition of 100% of Kestrel Fluid Controls Oy, a Finnish manufacturer of smart valve actuators "
        f"with FY2025 revenue of EUR {r.uniform(31, 38):.1f} million and 186 employees.",
        f"Indicative enterprise value: EUR {ev:.1f} million, equal to {r.uniform(9.5, 11.5):.1f}x FY2025 EBITDA. Financing from existing cash and a EUR 40m term loan."]},
             {"heading": "2. Synergies and timetable", "paragraphs": [
                 f"Revenue synergies of EUR {r.uniform(6, 9):.1f}m by 2028 from cross-selling actuators with HX pumps; cost synergies of EUR {r.uniform(1.8, 3):.1f}m from procurement.",
                 "Timetable: confirmatory due diligence until 15 March 2026, signing targeted for April 2026, closing subject to FCCA approval in Q3 2026.",
                 f"Deal team: {ceo.name} (sponsor), {cfo.name} (finance), {w.execs['gc'].name} (legal). Advisors: Fjell & Partners (M&A), Rautio Legal."]}]
    _pdf(w, "Project_Nordlys_Acquisition_Memo.pdf", "Project Nordlys - Acquisition Memo", "Strictly confidential - insider list",
         pages, ["Executive Office"], 4, ceo, "Confidential memo on the proposed acquisition of Kestrel Fluid Controls.", "M&A", date(2026, 1, 19))

    rows = []
    for k, p in w.execs.items():
        bonus = round(p.salary * r.uniform(.25, .6), -2)
        rows.append([p.name, p.title, p.salary, int(bonus), int(round(p.salary * r.uniform(.3, .9), -3)), r.choice(["Company car", "Car allowance"])])
    _csv(w, "Executive_Compensation_2025.csv", ["name", "title", "base_salary_eur", "bonus_paid_eur", "lti_grant_eur", "car_benefit"],
         rows, "Executive Compensation 2025", ["Executive Office", "Human Resources"], 4, w.execs["chro"],
         "Base salary, bonus and long-term incentive grants for the executive team.", "Compensation")

    cases = []
    for i in range(14):
        cases.append([f"LIT-{2023 + i % 3}-{100 + i}", r.choice(w.customers)["account_name"] if i % 2 else r.choice(w.suppliers),
                      r.choice(["Warranty claim", "Late delivery penalty", "Employment dispute", "Patent infringement (defendant)", "Contract termination", "Product liability"]),
                      r.choice(["Finland", "Germany", "France", "Spain", "Poland"]), r.randrange(40_000, 3_500_000, 5000),
                      r.choice(["Open", "Open", "Settlement negotiation", "Closed - settled", "Court of first instance"]),
                      r.choice(["Low", "Medium", "High"]), w.execs["gc"].name])
    _csv(w, "Litigation_Register.csv", ["case_id", "counterparty", "matter", "jurisdiction", "exposure_eur", "status", "risk", "owner"], cases,
         "Litigation Register", ["Legal & Compliance", "Executive Office"], 4, w.execs["gc"], "All open and recently closed legal disputes with exposure.", "Legal")


def finance_docs(w: World):
    r, cfo = w.rng, w.execs["cfo"]
    for q in range(1, 5):
        f = w.fin[q - 1]
        rev, cogs, opex, ebitda, da, ni = f["revenue"], f["cogs"], f["opex"], f["ebitda"], f["depreciation"], f["net_income"]
        pages = [{"heading": "1. Income statement", "paragraphs": [f"Unaudited consolidated results for Q{q} FY2025 (EUR million)."],
                  "table": [["Line item", f"Q{q} FY2025"], ["Revenue", f"{rev:.1f}"], ["Cost of goods sold", f"({cogs:.1f})"],
                            ["Gross profit", f"{rev - cogs:.1f}"], ["Operating expenses", f"({opex:.1f})"], ["EBITDA", f"{ebitda:.1f}"],
                            ["Depreciation & amortisation", f"({da:.1f})"], ["Net income", f"{ni:.1f}"]]},
                 {"heading": "2. Balance sheet highlights", "paragraphs": ["Balance at quarter end (EUR million)."],
                  "table": [["Item", "Amount"], ["Cash and equivalents", f"{f['cash']:.1f}"], ["Trade receivables", f"{f['receivables']:.1f}"],
                            ["Inventories", f"{f['inventories']:.1f}"], ["Trade payables", f"{f['payables']:.1f}"], ["Net debt", f"{f['net_debt']:.1f}"],
                            ["Equity", f"{f['equity']:.1f}"]]},
                 {"heading": "3. CFO commentary", "paragraphs": [
                     f"Gross margin was {f['gross_margin_pct']:.1f}%. DSO was {f['dso_days']} days and inventory turns {r.uniform(4.1, 5.2):.1f}. "
                     f"Free cash flow for the quarter was EUR {f['free_cash_flow']:.1f} million. - {cfo.name}, CFO"]}]
        _pdf(w, f"Financial_Statements_FY2025_Q{q}.pdf", f"Financial Statements Q{q} FY2025", "Group Finance - management accounts", pages,
             ["Finance"], 3, cfo, f"Q{q} FY2025 income statement, balance sheet and CFO commentary.", "Financial reporting",
             date(2025, 3 * q, 28) + timedelta(days=21))
    keys = list(w.fin[0].keys())
    _csv(w, "Financial_KPIs_FY2025.csv", keys, [[f[k] for k in keys] for f in w.fin], "Financial KPIs FY2025 (quarterly)",
         ["Finance", "Executive Office"], 3, cfo, "Quarterly revenue, margins, cash flow and balance-sheet KPIs (EUR million) - the source for finance dashboards.",
         "Financial reporting")

    rows = []
    cats = ["Personnel", "Travel", "IT & Software", "Facilities", "Training", "Consulting", "Marketing programs", "Equipment"]
    for d in DEPARTMENTS:
        for c in cats:
            base = r.randrange(20_000, 2_800_000, 1000) * (4 if c == "Personnel" else 1)
            rows.append([d, f"CC-{1000 + DEPARTMENTS.index(d) * 10}", c, base, int(base * r.uniform(.85, 1.12)), w.heads[d].name])
    _csv(w, "Budget_FY2026_by_Department.csv", ["department", "cost_center", "category", "budget_fy2026_eur", "actual_fy2025_eur", "budget_owner"],
         rows, "Budget FY2026 by Department", ["Finance"], 3, cfo, "Approved FY2026 operating budget per department and category vs FY2025 actuals.", "Budget")

    rows = []
    for i in range(420):
        inv = date(2026, 1, 1) + timedelta(days=r.randrange(0, 270))
        due = inv + timedelta(days=r.choice([30, 45, 60]))
        rows.append([f"INV-{77000 + i}", r.choice(w.suppliers), inv.isoformat(), due.isoformat(), r.randrange(800, 340_000, 10),
                     max(0, (date(2026, 10, 1) - due).days) if r.random() < .3 else 0, r.choice(["Approved", "Approved", "On hold - price variance", "Paid", "Paid", "Paid"])])
    _csv(w, "Accounts_Payable_Ledger_2026.csv", ["invoice_no", "supplier", "invoice_date", "due_date", "amount_eur", "days_overdue", "status"],
         rows, "Accounts Payable Ledger 2026", ["Finance"], 2, w.heads["Finance"], "Supplier invoices, due dates, overdue days and payment status.", "Accounts payable")

    rows = []
    for i in range(800):
        p = r.choice(w.staff)
        cat = r.choice(["Hotel", "Flight", "Meals", "Taxi", "Mileage", "Conference fee", "Client entertainment"])
        rows.append([f"EXP-{50000 + i}", p.name, p.department, cat, (date(2026, 1, 2) + timedelta(days=r.randrange(0, 260))).isoformat(),
                     round(r.uniform(12, 2400 if cat in ("Flight", "Conference fee") else 600), 2), r.choice(["Approved", "Approved", "Paid", "Rejected", "Pending"]),
                     w.heads[p.department].name])
    _csv(w, "Expense_Claims_2026.csv", ["claim_id", "employee", "department", "category", "date", "amount_eur", "status", "approver"], rows,
         "Expense Claims 2026", ["Finance", "Human Resources"], 2, w.heads["Finance"], "All 2026 employee expense claims with approver and status.", "Expenses")

    _docx(w, "Treasury_and_Signing_Authority_Policy.docx", "Treasury and Signing Authority Policy", [
        ("1. Signing limits", ["Payments and commitments require approval according to the limits below. Two signatures are required above EUR 50,000."],
         [["Role", "Single approval limit"], ["Department head", "EUR 50,000"], ["Executive (C-level)", "EUR 500,000"], ["CEO", "EUR 2,000,000"], ["Board", "Above EUR 2,000,000"]]),
        ("2. Cash management", ["Surplus cash is held with at least three banks rated A- or better, with no more than 40% at any single bank.",
                                "FX exposures above EUR 1 million equivalent are hedged at 75% for the following 12 months."], None),
        ("3. Bank accounts", [f"Opening or closing bank accounts requires approval by the CFO ({w.execs['cfo'].name}) and the CEO."], None)],
          ["Finance"], 2, cfo, "Signing limits, cash management and FX hedging rules.", "Policy")


def hr_docs(w: World):
    r, chro = w.rng, w.execs["chro"]
    rows = [[p.emp_id, p.name, p.title, p.department, p.site, p.email, p.phone, p.manager, p.start.isoformat()] for p in w.staff]
    _csv(w, "Staff_Directory.csv", ["employee_id", "name", "job_title", "department", "site", "email", "phone", "manager", "start_date"], rows,
         "Staff Directory (salaried staff)", ["Human Resources"], 1, w.heads["Human Resources"], "Contact details, role, site and manager of every salaried employee.", "Directory")

    bands = [["E1", "Executive", 210000, 300000, 420000], ["M1", "Head of function", 125000, 155000, 190000], ["P3", "Senior professional", 78000, 96000, 118000],
             ["P2", "Professional", 56000, 68000, 82000], ["P1", "Associate", 42000, 50000, 60000], ["S3", "Senior technician", 38000, 45000, 52000],
             ["S2", "Technician / operator", 31000, 36000, 42000]]
    _csv(w, "Salary_Bands_2026.csv", ["grade", "level_description", "min_eur", "midpoint_eur", "max_eur"], bands, "Salary Bands 2026",
         ["Human Resources"], 3, chro, "Annual base salary ranges per grade for 2026.", "Compensation")

    by_dept = {}
    for p in w.staff:
        by_dept[p.department] = by_dept.get(p.department, 0) + 1
    table = [["Department", "Salaried headcount", "Hires 2025", "Leavers 2025", "Attrition"]]
    for d, n in sorted(by_dept.items()):
        hires, leav = r.randint(2, max(3, n // 5)), r.randint(1, max(2, n // 7))
        table.append([d, str(n), str(hires), str(leav), f"{100 * leav / n:.1f}%"])
    _pdf(w, "Headcount_and_Attrition_Report_2025.pdf", "Headcount and Attrition Report 2025", "People Analytics",
         [{"heading": "1. Summary", "paragraphs": [f"Salaried headcount at 31 December 2025 was {len(w.staff)}. Group voluntary attrition was 9.8%, "
                                                   "above the 7% target. Engineering and Sales had the highest regretted attrition."]},
          {"heading": "2. By department", "paragraphs": ["Salaried staff only; hourly production staff are reported in the Operations report."], "table": table}],
         ["Human Resources"], 3, chro, "Headcount, hires, leavers and attrition per department.", "People analytics")

    _pdf(w, "Benefits_Guide_2026.pdf", "Employee Benefits Guide 2026", "Human Resources", [
        {"heading": "1. Health", "paragraphs": ["All employees are covered by private occupational health care from day one, including dental check-ups once a year.",
                                                "Eyeglasses for screen work are reimbursed up to EUR 250 every two years."]},
        {"heading": "2. Wellbeing", "paragraphs": ["The annual wellbeing allowance is EUR 400, usable for sports, culture and massage.",
                                                   "Employees may take one paid volunteering day per year."]},
        {"heading": "3. Pension and insurance", "paragraphs": ["The company contributes 4% of base salary to the voluntary pension plan when the employee contributes at least 2%.",
                                                              "Travel insurance covers all business trips and the employee's spouse when travelling together."]},
        {"heading": "4. Bike and commuting", "paragraphs": ["The bike benefit allows leasing a bike of up to EUR 3,000 through salary deduction.",
                                                            "Public transport tickets are subsidised by 50% up to EUR 60 per month."]}],
         ["Human Resources"], 1, w.heads["Human Resources"], "Health, wellbeing, pension and commuting benefits for all staff.", "Benefits")

    rows = []
    for i in range(46):
        d = r.choice(DEPARTMENTS[1:])
        rows.append([f"REQ-{3100 + i}", r.choice(TITLES[d]), d, r.choice(SITE_NAMES), w.heads[d].name,
                     (date(2026, 6, 1) + timedelta(days=r.randrange(0, 120))).isoformat(), r.choice(["Open", "Open", "Interviewing", "Offer extended"]),
                     r.choice(["Permanent", "Permanent", "Fixed-term 12 months"])])
    _csv(w, "Open_Positions.csv", ["requisition_id", "job_title", "department", "site", "hiring_manager", "posted_date", "status", "contract"], rows,
         "Open Positions", ["Human Resources"], 1, w.heads["Human Resources"], "Current vacancies, hiring managers and recruitment status.", "Recruitment")

    courses = [("Lean Six Sigma Green Belt", 5, "Operations"), ("Leading High-Performing Teams", 3, "All managers"), ("GDPR Essentials", 0.5, "All staff"),
               ("Advanced CFD with OpenFOAM", 4, "Engineering"), ("Negotiation Skills", 2, "Sales, Procurement"), ("First Aid at Work", 1, "Operations"),
               ("Cybersecurity Awareness", 0.5, "All staff"), ("Finance for Non-Financial Managers", 2, "All managers"),
               ("IEC 61508 Functional Safety", 3, "Engineering"), ("Forklift Licence Renewal", 1, "Operations")]
    md = ["# Training Catalogue 2026", "", "Courses are booked via the learning portal. Cost is charged to the department budget.", ""]
    for i, (c, days, aud) in enumerate(courses, 1):
        md += [f"## {c}", "", f"- Duration: {days} day(s)", f"- Audience: {aud}", f"- Cost per participant: EUR {r.randrange(80, 2400, 10)}",
               f"- Provider: {r.choice(['Internal academy', 'Tampere University of Applied Sciences', 'Nordic Lean Institute', 'SecureMind Training'])}", ""]
    _md(w, "Training_Catalogue_2026.md", "\n".join(md), "Training Catalogue 2026", ["Human Resources"], 1, w.heads["Human Resources"],
        "Available training courses, duration, audience, cost and provider.", "Learning")

    rows = []
    for p in r.sample(w.staff, 240):
        rows.append([p.emp_id, p.name, p.department, p.manager, r.choices(["1 - Below", "2 - Partly meets", "3 - Meets", "4 - Exceeds", "5 - Outstanding"], [2, 8, 55, 28, 7])[0],
                     r.choice(["Low", "Medium", "High"]), r.choice(["No", "No", "No", "Yes"]), f"{r.choice([0, 2, 2.5, 3, 3, 3.5, 4, 5])}%"])
    _csv(w, "Performance_Calibration_2025.csv", ["employee_id", "name", "department", "manager", "rating", "potential", "flight_risk", "merit_increase"], rows,
         "Performance Calibration 2025", ["Human Resources", "Executive Office"], 4, chro, "Calibrated ratings, potential, flight risk and merit increases.", "Performance")


def engineering_docs(w: World):
    r, cto = w.rng, w.execs["cto"]
    secs = []
    for i, (name, q, desc) in enumerate([
        ("SmartFlow 2.0 platform", "Q2 2026", "cloud analytics for pump fleets with predictive maintenance alerts"),
        ("HX-700 high-pressure pump", "Q4 2026", "new 700-series for desalination with 160 m maximum head"),
        ("CU-5 controller", "Q1 2027", "next-generation controller with edge AI and IEC 62443 certification"),
        ("IE5 motor upgrade across HX range", "Q3 2026", "ultra-premium efficiency motors on all HX models above 15 kW"),
        ("Kestrel actuator integration", "Q2 2027", "bundle smart valve actuators with HX pumps (subject to Project Nordlys)")]):
        secs.append((f"{i + 1}. {name}", [f"Planned release: {q}. Scope: {desc}.",
                                          f"Product owner: {r.choice([p for p in w.staff if p.department == 'Engineering']).name}. "
                                          f"Development budget: EUR {r.uniform(.8, 4.5):.1f} million. Status: {r.choice(['On track', 'At risk - resourcing', 'On track', 'Design phase'])}."], None))
    _docx(w, "Product_Roadmap_2026-2027.docx", "Product Roadmap 2026-2027", secs, ["Engineering", "Sales", "Marketing"], 3, cto,
          "Planned product releases, owners, budgets and status.", "Roadmap")

    for model in ["HX-200", "HX-280", "HX-360", "HX-450", "HX-520", "HX-600"]:
        eng = r.choice([p for p in w.staff if p.department == "Engineering"])
        rows = [["Test", "Requirement", "Result", "Verdict"]]
        for t, req in [("Hydrostatic pressure", "1.5x rated, 10 min"), ("Vibration at BEP", "< 3.5 mm/s RMS"), ("Noise at 1 m", "< 78 dB(A)"),
                       ("Efficiency at BEP", ">= 78%"), ("Endurance run", "1,000 h"), ("Seal leakage", "< 5 ml/h")]:
            res = {"Hydrostatic pressure": "No leakage", "Vibration at BEP": f"{r.uniform(1.4, 3.8):.1f} mm/s", "Noise at 1 m": f"{r.uniform(68, 80):.0f} dB(A)",
                   "Efficiency at BEP": f"{r.uniform(74, 86):.1f}%", "Endurance run": f"{r.choice([1000, 1000, 1000, 840])} h", "Seal leakage": f"{r.uniform(.5, 6):.1f} ml/h"}[t]
            rows.append([t, req, res, r.choice(["Pass", "Pass", "Pass", "Pass", "Fail"])])
        _docx(w, f"Type_Test_Report_{model}.docx", f"Type Test Report {model}", [
            ("1. Summary", [f"Type testing of the {model} was performed at the Lyon Systems test lab between {r.randint(1, 20)} and {r.randint(21, 28)} {r.choice(['March', 'April', 'May', 'June'])} 2026.",
                            f"Test engineer: {eng.name}. Witnessed by Quality: {r.choice([p for p in w.staff if p.department == 'Quality']).name}."], None),
            ("2. Results", ["Results against the product requirement specification:"], rows)],
              ["Engineering", "Quality"], 2, eng, f"Type test results for the {model}: pressure, vibration, noise, efficiency and endurance.", "Test report")

    md = ["# Architecture Decision Records - SmartFlow", ""]
    for i, (t, dec) in enumerate([("Message broker", "Use MQTT with a Mosquitto cluster for pump telemetry"), ("Time-series store", "Use TimescaleDB instead of InfluxDB"),
                                  ("Edge runtime", "Run analytics on the CU-4 Pro using containerised Python"), ("Authentication", "Use OpenID Connect with the corporate identity provider"),
                                  ("Data retention", "Keep raw telemetry for 13 months and hourly aggregates for 7 years"), ("Frontend", "Build the customer portal in React with TypeScript")], 1):
        md += [f"## ADR-{i:03d}: {t}", "", f"- Status: Accepted", f"- Date: {(date(2025, 9, 1) + timedelta(days=30 * i)).isoformat()}",
               f"- Decision: {dec}.", f"- Decided by: {r.choice([p for p in w.staff if p.department == 'Engineering']).name}", ""]
    _md(w, "SmartFlow_Architecture_Decisions.md", "\n".join(md), "SmartFlow Architecture Decision Records", ["Engineering", "IT"], 2, cto,
        "Key technical decisions for the SmartFlow platform.", "Architecture")

    rows = []
    for i in range(160):
        rows.append([f"ECO-{6000 + i}", r.choice(PRODUCTS[:14]), r.choice(["Change impeller material to duplex steel", "Update seal supplier", "Firmware parameter change",
                                                                         "Revise drawing tolerance", "Replace obsolete capacitor", "Add corrosion coating", "Change bearing grade"]),
                     r.choice(["Draft", "In review", "Approved", "Implemented", "Rejected"]), r.choice([p for p in w.staff if p.department == "Engineering"]).name,
                     (date(2025, 10, 1) + timedelta(days=r.randrange(0, 330))).isoformat(), r.choice(["Low", "Medium", "High"])])
    _csv(w, "Engineering_Change_Orders.csv", ["eco_id", "product", "description", "status", "requested_by", "date", "impact"], rows,
         "Engineering Change Orders", ["Engineering", "Quality", "Operations"], 2, cto, "Engineering change orders with product, status and impact.", "Change control")


def operations_docs(w: World):
    r, coo = w.rng, w.execs["coo"]
    rows = []
    parts = ["Impeller", "Pump casing", "Mechanical seal", "Shaft", "Bearing 6309", "Motor 15 kW", "Motor 45 kW", "Coupling", "Gasket set",
             "CU-4 main board", "Pressure sensor", "Cable gland", "Paint (litres)", "Bolt M16"]
    for i in range(400):
        part = r.choice(parts)
        rows.append([f"SKU-{40000 + i}", f"{part} {r.choice(PRODUCTS[:12])}", r.choice(SITE_NAMES), r.randint(0, 2400), r.randint(20, 400),
                     round(r.uniform(.4, 1800), 2), r.choice(w.suppliers), r.randint(5, 90)])
    _csv(w, "Inventory_Levels.csv", ["sku", "description", "site", "on_hand_qty", "reorder_point", "unit_cost_eur", "preferred_supplier", "lead_time_days"],
         rows, "Inventory Levels", ["Operations", "Procurement"], 2, w.heads["Operations"], "Stock on hand, reorder points and supplier per SKU and site.", "Inventory")

    rows = []
    for i in range(260):
        rows.append([f"PO-PLAN-{i + 1:04d}", r.choice(PRODUCTS[:12]), r.choice(SITE_NAMES[:7]), f"2026-W{r.randint(1, 13):02d}", r.randint(5, 220),
                     r.choice(w.customers)["account_name"], r.choice(["Planned", "Released", "In production", "Completed"])])
    _csv(w, "Production_Plan_2026_Q1.csv", ["order_id", "product", "site", "week", "quantity", "customer", "status"], rows,
         "Production Plan 2026 Q1", ["Operations"], 2, w.heads["Operations"], "Weekly production orders per site and product for Q1 2026.", "Planning")

    rows = []
    for i in range(120):
        rows.append([f"SI-2025-{i + 1:03d}", (date(2025, 1, 3) + timedelta(days=r.randrange(0, 360))).isoformat(), r.choice(SITE_NAMES),
                     r.choice(["Slip/trip", "Hand injury", "Forklift near miss", "Chemical splash", "Manual handling", "Burn", "Falling object"]),
                     r.choices(["Near miss", "First aid", "Medical treatment", "Lost time"], [5, 4, 2, 1])[0], r.randint(0, 14),
                     r.choice(["Closed", "Closed", "Action open"]), r.choice([p for p in w.staff if p.department == "Operations"]).name])
    _csv(w, "Safety_Incident_Log_2025.csv", ["incident_id", "date", "site", "type", "severity", "days_lost", "status", "investigator"], rows,
         "Safety Incident Log 2025", ["Operations", "Quality"], 2, w.heads["Operations"], "All 2025 safety incidents and near misses by site.", "EHS")


def sales_marketing_docs(w: World):
    r, cco = w.rng, w.execs["cco"]
    head = list(w.customers[0].keys())
    _csv(w, "Customer_Accounts.csv", head, [list(c.values()) for c in w.customers], "Customer Accounts", ["Sales"], 2, w.heads["Sales"],
         "All customer accounts with segment, annual revenue, account manager and payment terms.", "CRM")

    rows = []
    stages = [("Qualification", 10), ("Needs analysis", 25), ("Proposal", 45), ("Negotiation", 70), ("Closed won", 100), ("Closed lost", 0)]
    for i in range(540):
        c = r.choice(w.customers)
        st, pr = r.choice(stages)
        rows.append([f"OPP-{90000 + i}", c["account_name"], r.choice(PRODUCTS), st, r.randrange(8_000, 1_900_000, 500), pr,
                     (date(2026, 1, 15) + timedelta(days=r.randrange(0, 420))).isoformat(), c["account_manager"], c["country"]])
    _csv(w, "Sales_Pipeline_2026.csv", ["opportunity_id", "account", "product", "stage", "value_eur", "probability_pct", "expected_close", "owner", "country"],
         rows, "Sales Pipeline 2026", ["Sales"], 2, w.heads["Sales"], "Open and closed opportunities with stage, value, probability and owner.", "CRM")

    rows = []
    for p in PRODUCTS:
        lst = r.randrange(900, 48_000, 10) if p.startswith("HX") else r.randrange(150, 12_000, 10)
        rows.append([p, lst, f"{r.choice([5, 8, 10, 12, 15])}%", f"{r.choice([18, 22, 25, 30])}%", r.choice(["2 years", "3 years", "1 year"])])
    _csv(w, "Price_List_2026.csv", ["product", "list_price_eur", "max_discount_account_manager", "max_discount_head_of_sales", "warranty"], rows,
         "Price List 2026", ["Sales", "Marketing"], 2, w.heads["Sales"], "2026 list prices, discount authority and warranty per product.", "Pricing")

    regions = ["Nordics", "DACH", "Iberia", "France & Benelux", "Central & Eastern Europe"]
    table = [["Region", "Revenue FY2025", "Growth", "Win rate", "Regional lead"]]
    for g in regions:
        table.append([g, f"EUR {r.uniform(48, 120):.1f}m", f"{r.uniform(-3, 14):.1f}%", f"{r.uniform(22, 41):.0f}%", r.choice([p for p in w.staff if p.department == "Sales"]).name])
    _pdf(w, "Sales_Performance_Review_FY2025.pdf", "Sales Performance Review FY2025", "Commercial Excellence", [
        {"heading": "1. Highlights", "paragraphs": [f"FY2025 order intake reached EUR {sum(f['order_intake'] for f in w.fin):.1f} million. Service revenue grew {r.uniform(9, 16):.1f}% and now represents "
                                                    f"{r.uniform(17, 21):.1f}% of sales. The largest single order was EUR {r.uniform(4, 9):.1f} million from {r.choice(w.customers)['account_name']}."]},
        {"heading": "2. Regional performance", "paragraphs": ["Revenue by sales region:"], "table": table},
        {"heading": "3. FY2026 priorities", "paragraphs": [f"Raise win rate in DACH, launch SmartFlow 2.0 bundles and open the Stockholm office. - {cco.name}, CCO"]}],
         ["Sales"], 3, cco, "FY2025 sales results by region, win rates and FY2026 priorities.", "Sales reporting")

    # monthly order bookings by region, scaled so calendar 2025 equals FY2025 order intake
    reg_share = {"Nordics": .27, "DACH": .24, "Iberia": .15, "France & Benelux": .19, "Central & Eastern Europe": .15}
    months = [(y, m_) for y in (2024, 2025, 2026) for m_ in range(1, 13) if (y, m_) <= (2026, 9)]
    season = {1: .86, 2: .9, 3: 1.08, 4: .97, 5: 1.02, 6: 1.12, 7: .78, 8: .84, 9: 1.06, 10: 1.05, 11: 1.08, 12: 1.24}
    raw = {(y, m_, g): season[m_] * (1 + .055 * (y - 2025) + .004 * m_) * sh * r.uniform(.9, 1.1)
           for (y, m_) in months for g, sh in reg_share.items()}
    scale = sum(f["order_intake"] for f in w.fin) / sum(v for (y, _, _), v in raw.items() if y == 2025)
    rows = []
    for (y, m_, g), v in raw.items():
        orders = round(v * scale * 1e6, -3)
        quotes = int(orders / r.uniform(38_000, 52_000))
        rows.append([f"{y}-{m_:02d}", g, int(orders), quotes, round(r.uniform(24, 41), 1)])
    _csv(w, "Order_Bookings_Monthly.csv", ["month", "region", "orders_eur", "quotes_issued", "win_rate_pct"], rows,
         "Order Bookings by Month and Region", ["Sales", "Executive Office"], 3, cco,
         "Monthly order intake, quotes and win rate per sales region, January 2024 to September 2026.", "Sales reporting")

    rows = []
    for i in range(48):
        rows.append([f"CMP-{i + 1:03d}", r.choice(["Webinar", "Trade fair", "LinkedIn ads", "Email nurture", "Customer event", "Google Ads"]),
                     r.choice(["HX range", "SmartFlow", "CU-4", "Service agreements"]), r.choice(regions), r.randrange(2_000, 140_000, 500),
                     r.randint(20, 1400), r.randint(0, 60), f"{r.uniform(.5, 9):.1f}x", (date(2025, 1, 10) + timedelta(days=r.randrange(0, 340))).isoformat()])
    _csv(w, "Marketing_Campaigns_2025.csv", ["campaign_id", "channel", "product_focus", "region", "spend_eur", "leads", "opportunities", "roi", "start_date"],
         rows, "Marketing Campaigns 2025", ["Marketing", "Sales"], 2, w.heads["Marketing"], "Campaign spend, leads, opportunities and ROI.", "Marketing")

    comps = [("Grundfors Pumps", "market leader, strong brand, 12% price premium"), ("Wilmar Fluid", "aggressive pricing in Iberia, weaker service network"),
             ("KSV Pumpen", "strong in DACH chemicals, long lead times"), ("Sulzbach Flow", "premium high-pressure range, competing with HX-700")]
    _docx(w, "Competitor_Analysis_2026.docx", "Competitor Analysis 2026", [
        (f"{i + 1}. {n}", [f"Positioning: {d}.", f"Estimated European revenue EUR {r.randint(150, 2400)}m. Our win rate against them in 2025: {r.randint(18, 55)}%."], None)
        for i, (n, d) in enumerate(comps)], ["Sales", "Marketing"], 3, w.heads["Marketing"], "Positioning, size and win rates against main competitors.", "Market intelligence")


def legal_it_procurement_docs(w: World):
    r = w.rng
    rows = []
    for i in range(150):
        cp = r.choice([c["account_name"] for c in w.customers] + w.suppliers)
        start = date(2022, 1, 1) + timedelta(days=r.randrange(0, 1400))
        rows.append([f"CON-{8000 + i}", cp, r.choice(["Framework supply agreement", "NDA", "Service agreement", "Distribution agreement", "Software licence", "Lease"]),
                     start.isoformat(), (start + timedelta(days=365 * r.choice([1, 2, 3, 5]))).isoformat(), r.randrange(0, 4_000_000, 1000),
                     r.choice(["Yes", "No"]), r.choice([p for p in w.staff if p.department == "Legal & Compliance"]).name])
    _csv(w, "Contract_Register.csv", ["contract_id", "counterparty", "type", "start_date", "end_date", "value_eur", "auto_renewal", "contract_owner"], rows,
         "Contract Register", ["Legal & Compliance", "Procurement"], 2, w.heads["Legal & Compliance"], "Active contracts with counterparty, value, term and owner.", "Contracts")

    _pdf(w, "Data_Protection_Policy.pdf", "Data Protection Policy", "Legal & Compliance", [
        {"heading": "1. Principles", "paragraphs": ["Personal data is processed lawfully, for specified purposes and kept no longer than necessary, in line with the GDPR.",
                                                    f"The Data Protection Officer is {r.choice([p for p in w.staff if p.department == 'Legal & Compliance']).name}, reachable at privacy@halvorsen-is.example."]},
        {"heading": "2. Breaches", "paragraphs": ["Personal data breaches must be reported to the DPO within 24 hours of discovery; the DPO notifies the authority within 72 hours where required."]},
        {"heading": "3. Data subject requests", "paragraphs": ["Access and deletion requests are answered within 30 days. Requests are logged in the privacy request register."]}],
         ["Legal & Compliance"], 1, w.heads["Legal & Compliance"], "GDPR principles, breach reporting and data subject requests.", "Policy")

    _md(w, "Anti_Bribery_and_Gifts_Policy.md", "\n".join([
        "# Anti-Bribery and Gifts Policy", "", "## Gifts and hospitality", "",
        "Gifts above EUR 75 must not be accepted. Hospitality above EUR 150 per person requires prior approval by the Compliance Officer.", "",
        "## Third parties", "", "Agents and distributors are screened before appointment and every two years thereafter.", "",
        "## Reporting", "", "Concerns can be raised anonymously through the whistleblowing hotline, available 24 hours a day."]),
        "Anti-Bribery and Gifts Policy", ["Legal & Compliance"], 1, w.heads["Legal & Compliance"], "Rules on gifts, hospitality, third parties and reporting.", "Policy")

    rows = []
    for i in range(700):
        p = r.choice(w.staff)
        kind = r.choice(["Laptop", "Laptop", "Monitor", "Mobile phone", "Docking station", "Tablet"])
        model = {"Laptop": r.choice(["Lenovo ThinkPad T14 Gen 5", "Dell Latitude 7450", "HP EliteBook 840 G11"]), "Monitor": "Dell P2725H",
                 "Mobile phone": r.choice(["iPhone 15", "Samsung Galaxy S24"]), "Docking station": "Lenovo USB-C Dock", "Tablet": "iPad Air"}[kind]
        purch = date(2022, 1, 1) + timedelta(days=r.randrange(0, 1500))
        rows.append([f"IT-{100000 + i}", kind, model, f"SN{r.randint(10**7, 10**8 - 1)}", p.name, p.department, purch.isoformat(),
                     (purch + timedelta(days=365 * 3)).isoformat(), r.choice(["In use", "In use", "In use", "In repair", "Retired"])])
    _csv(w, "IT_Asset_Inventory.csv", ["asset_tag", "type", "model", "serial_number", "assigned_to", "department", "purchase_date", "refresh_due", "status"],
         rows, "IT Asset Inventory", ["IT"], 2, w.heads["IT"], "Laptops, phones and peripherals with assignee and refresh dates.", "IT assets")

    systems = ["SAP S/4HANA", "Workday", "Salesforce", "SmartFlow admin", "Teamcenter PLM", "Jira", "Power BI", "Treasury (Kyriba)", "Payroll", "MES"]
    rows = []
    for s in systems:
        for d in DEPARTMENTS:
            rows.append([s, d, r.choice(["None", "Read", "Read", "Read/Write", "Admin"]) if d != "IT" else "Admin", w.heads[d].name])
    _csv(w, "System_Access_Matrix.csv", ["system", "department", "access_level", "approver"], rows, "System Access Matrix", ["IT", "Legal & Compliance"], 3,
         w.execs["cio"], "Which department has what access to each business system, and who approves it.", "Access control")

    _docx(w, "IT_Security_Standard.docx", "IT Security Standard", [
        ("1. Identity", ["Multi-factor authentication is required for all accounts. Privileged accounts use hardware security keys.",
                         "Access rights are reviewed every 6 months by department heads."], None),
        ("2. Devices", ["Laptops use full-disk encryption and automatic screen lock after 5 minutes. USB mass storage is blocked by default."], None),
        ("3. Incident response", ["Security incidents are reported to the Security Operations Centre within 2 hours via the IT service desk or +358 20 555 0100."], None)],
          ["IT"], 1, w.execs["cio"], "Mandatory security controls for identities, devices and incident response.", "Policy")

    rows = []
    for i in range(900):
        s = r.choice(w.suppliers)
        rows.append([f"PO-{450000 + i}", s, r.choice(["Castings", "Motors", "Seals", "Bearings", "Electronics", "Logistics", "MRO", "IT hardware", "Consulting"]),
                     (date(2026, 1, 2) + timedelta(days=r.randrange(0, 270))).isoformat(), r.randrange(300, 480_000, 10),
                     r.choice([p for p in w.staff if p.department == "Procurement"]).name, r.choice(["Open", "Received", "Received", "Invoiced", "Cancelled"]),
                     r.choice(SITE_NAMES)])
    _csv(w, "Purchase_Orders_2026.csv", ["po_number", "supplier", "category", "order_date", "value_eur", "buyer", "status", "delivery_site"], rows,
         "Purchase Orders 2026", ["Procurement", "Finance"], 2, w.heads["Procurement"], "All 2026 purchase orders with supplier, value, buyer and status.", "Purchasing")


def company_docs(w: World):
    ceo = w.execs["ceo"]
    lines = ["# Company Overview", "", f"{COMPANY} (HIS) designs and manufactures centrifugal pumps, pump controllers and digital monitoring services "
             "for water, energy and industrial customers across Europe.", "", "## Key facts", "",
             "- Founded: 1962 in Tampere, Finland", "- Group revenue FY2025: EUR 412.6 million", "- Employees: about 3,940 (including production staff)",
             "- Manufacturing sites: " + ", ".join(SITE_NAMES), f"- Chief Executive Officer: {ceo.name}", "", "## Executive team", ""]
    lines += [f"- {p.title}: {p.name}" for p in w.execs.values()]
    lines += ["", "## Department heads", ""] + [f"- {d}: {h.name} ({h.site})" for d, h in w.heads.items()]
    _md(w, "Company_Overview.md", "\n".join(lines), "Company Overview", ["Executive Office"], 1, ceo, "Who we are, key facts, executive team and department heads.", "Company")

    town = ["# CEO Town Hall - Q4 2025 summary", "", f"Presented by {ceo.name} on 11 December 2025.", "", "## Results", "",
            "FY2025 revenue grew 6.3% to EUR 412.6 million and safety performance improved to an LTIFR of 0.62.", "",
            "## What's next", "", "Our 2026 priorities are SmartFlow 2.0, the HX-700 launch and reducing attrition. "
            "Every employee will receive a one-off recognition payment of EUR 500 in January 2026.", "",
            "## Q&A highlights", "", "Asked about remote work, the CEO confirmed the 3-days-per-week policy remains unchanged in 2026."]
    _md(w, "CEO_Town_Hall_Q4_2025.md", "\n".join(town), "CEO Town Hall Q4 2025", ["Executive Office"], 1, ceo, "Summary of the Q4 2025 all-hands meeting.", "Communications")

    rows = [[s, SITE_COUNTRY[s], w.rng.choice(["Manufacturing", "Assembly", "Foundry", "Machining", "R&D and manufacturing", "Service"]),
             f"{w.rng.randint(1, 99)} Industrial Road", next((h.name for h in w.heads.values() if h.site == s), w.execs["coo"].name),
             f"+{w.rng.randint(30, 49)} {w.rng.randint(10, 99)} {w.rng.randint(100, 999)} {w.rng.randint(1000, 9999)}"] for s in SITE_NAMES]
    _csv(w, "Office_and_Site_Locations.csv", ["site", "country", "function", "address", "site_contact", "reception_phone"], rows,
         "Office and Site Locations", ["Operations"], 1, w.execs["coo"], "All sites with country, function, address and contact.", "Company")


# Existing assignment corpus, re-used with access levels.
LEGACY = {
    "HIS_Employee_Handbook_2026.pdf": ("Employee Handbook 2026", ["Human Resources"], 1, "Policy"),
    "HIS_Annual_Operations_Report_FY2025.pdf": ("Annual Operations Report FY2025", ["Operations"], 3, "Operations reporting"),
    "HIS_Quality_Compliance_Manual.pdf": ("Quality and Compliance Manual", ["Quality"], 1, "Quality"),
    "HIS_Site_Energy_and_Sustainability_Review_2025.pdf": ("Site Energy and Sustainability Review 2025", ["Operations"], 1, "Sustainability"),
    "HX_Pump_Series_Technical_Specification.docx": ("HX Pump Series Technical Specification", ["Engineering"], 1, "Product"),
    "HX_Pump_Selection_Guide.md": ("HX Pump Selection Guide", ["Engineering", "Sales"], 1, "Product"),
    "CU4_Controller_Firmware_Release_Notes.docx": ("CU-4 Firmware Release Notes", ["Engineering"], 2, "Product"),
    "CU4_Known_Issues.md": ("CU-4 Known Issues", ["Engineering"], 2, "Product"),
    "Supplier_Audit_Report_2025.pdf": ("Supplier Audit Report 2025", ["Procurement", "Quality"], 2, "Supplier quality"),
    "supplier_register.csv": ("Supplier Register", ["Procurement"], 2, "Suppliers"),
    "equipment_maintenance_schedule.csv": ("Equipment Maintenance Schedule", ["Operations"], 2, "Maintenance"),
    "maintenance_work_orders_2026.csv": ("Maintenance Work Orders 2026", ["Operations"], 2, "Maintenance"),
    "incident_postmortems.md": ("IT Incident Post-Mortems 2025", ["IT"], 2, "IT operations"),
    "engineering_meeting_notes.md": ("Engineering Leadership Meeting Notes", ["Engineering"], 2, "Meetings"),
}


def build(seed: int = 2026) -> dict:
    if DOC_DIR.exists():
        shutil.rmtree(DOC_DIR)
    DOC_DIR.mkdir(parents=True)
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    w = build_world(seed)
    for fn in (executive_docs, finance_docs, hr_docs, engineering_docs, operations_docs, sales_marketing_docs,
               legal_it_procurement_docs, company_docs):
        fn(w)
    for name, (title, depts, level, cat) in LEGACY.items():
        src = config.CORPUS_DIR / name
        if src.exists():
            shutil.copy(src, DOC_DIR / name)
            owner = w.heads[depts[0]]
            _add(w, name, title, depts, level, owner, f"{title}.", cat, date(2026, 2, 1))
    MANIFEST.write_text(json.dumps([asdict(d) for d in w.docs], indent=2))
    _write_accounts(w)
    return {"documents": len(w.docs), "staff": len(w.staff), "customers": len(w.customers),
            "by_level": {LEVELS[k]: sum(d.level == k for d in w.docs) for k in LEVELS}}


# ---------------------------------------------------------------------------
# demo portal accounts (test credentials for this local demo only)
# ---------------------------------------------------------------------------
def _write_accounts(w: World):
    rows = []
    for key, p in w.execs.items():
        rows.append([key, f"{key.capitalize()}@2026", p.name, p.title, p.department, "executive"])
    slug = {"Executive Office": "office", "Finance": "finance", "Human Resources": "hr", "Engineering": "engineering", "Operations": "operations",
            "Sales": "sales", "Marketing": "marketing", "Legal & Compliance": "legal", "IT": "it", "Quality": "quality", "Procurement": "procurement"}
    for d, h in w.heads.items():
        rows.append([f"head.{slug[d]}", f"Head.{slug[d].capitalize()}@2026", h.name, h.title, d, "manager"])
    for d in DEPARTMENTS[1:]:
        p = next(x for x in w.staff if x.department == d and x not in w.heads.values() and x not in w.execs.values())
        rows.append([slug[d], f"{slug[d].capitalize()}@2026", p.name, p.title, d, "employee"])
    it = next(x for x in w.staff if x.department == "IT" and "Administrator" in x.title)
    rows.append(["admin", "Admin@2026", it.name, "Portal Administrator", "IT", "admin"])
    with open(ACCOUNTS, "w", newline="") as f:
        wr = csv.writer(f)
        wr.writerow(["username", "password", "full_name", "title", "department", "role"])
        wr.writerows(rows)
